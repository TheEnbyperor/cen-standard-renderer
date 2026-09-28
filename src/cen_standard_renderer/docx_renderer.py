import pathlib
import typing
import re
import docx
import docx.document
import docx.opc.constants
import docx.oxml
import docx.oxml.ns
import docx.shared
import docx.enum.text
import docx.enum.table
import docx.enum.style
from . import ir, standard


class DocxRenderer:
    def __init__(self, template: pathlib.Path | None = None, asn1 = None) -> None:
        self.doc = docx.Document(str(template)) if template is not None else docx.Document()
        self.bookmark_id = 1
        self.asn1 = asn1
        self.asn1_bookmark_id = 1
        self.asn1_bookmarks = {}

    def render(self, src: standard.StandardSource, output: pathlib.Path) -> None:
        self._render_front_matter(src)
        self._render_introduction(src.introduction)
        self._render_scope(src.scope)
        self._render_blocks(src.terms, self.doc)
        for sec in src.sections:
            p = self.doc.add_paragraph(style="Heading 1")
            p.add_run(sec.metadata.get("title"))
            self._render_blocks(sec.blocks, self.doc)
        self.doc.save(str(output.absolute()))

    def _render_front_matter(self, src: standard.StandardSource, ) -> None:
        title = src.metadata.get("title")
        number = src.metadata.get("document_number")

        if title is not None:
            self.doc.add_paragraph(str(title), style="zzCover")
        if number is not None:
            self.doc.add_paragraph(str(number), style="zzCover")

        if title or number:
            self.doc.add_paragraph().add_run().add_break(docx.enum.text.WD_BREAK.PAGE)

        p = self.doc.add_paragraph(style="zzContents")
        p.add_run("Contents")
        p = self.doc.add_paragraph()
        run = p.add_run()
        add_field(run, r'TOC \o "1-4" \h \z \u', "Right click to update ToC")
        run.add_break(docx.enum.text.WD_BREAK.PAGE)

    def _render_blocks(self, blocks: typing.Iterable[ir.Block], base) -> None:
        for block in blocks:
            if isinstance(block, ir.Heading):
                self._heading(block, base)
            elif isinstance(block, ir.Paragraph):
                p = base.add_paragraph(style="Body Text")
                self._inlines(p, block.children)
            elif isinstance(block, ir.CodeBlock):
                p = base.add_paragraph(style="Code (-)")
                p.add_run(block.text)
            elif isinstance(block, ir.BulletList):
                self._list(block.items, False, base)
            elif isinstance(block, ir.OrderedList):
                self._list(block.items, True, base)
            elif isinstance(block, ir.ASN1Type):
                self._asn1_type(block, base)
            elif isinstance(block, ir.SpecialBlock):
                continue
            else:
                raise NotImplementedError(type(block).__name__)

    def _render_introduction(self, blocks: typing.Iterable[ir.Block]) -> None:
        p = self.doc.add_paragraph(style="Intro Title")
        p.add_run("Introduction")
        self._render_blocks(blocks, self.doc)
        self.doc.add_paragraph().add_run().add_break(docx.enum.text.WD_BREAK.PAGE)

    def _render_scope(self, blocks: typing.Iterable[ir.Block]) -> None:
        p = self.doc.add_paragraph(style="Heading 1")
        p.add_run("Scope")
        self._render_blocks(blocks, self.doc)

    def _heading(self, heading: ir.Heading, base) -> None:
        if heading.level == ir.HEADING_LEVEL_TERM:
            style = "Term(s)"
        else:
            style = f"Heading {min(heading.level + 1, 6)}"
        p = base.add_paragraph(style=style)
        self._inlines(p, heading.children)

        if heading.identifier:
            add_bookmark(p, heading.identifier, self.bookmark_id)
            self.bookmark_id += 1

    def _new_numbering(self, paragraph, start: int = 1) -> int:
        style = paragraph.style
        old_num_id = int(style.element.pPr.numPr.numId.val)
        doc_numbering = self.doc.part.numbering_part.element
        old_num = next(n for n in doc_numbering.num_lst if n.numId == old_num_id)
        new_num = doc_numbering.add_num(old_num.abstractNumId.val)
        new_num.add_lvlOverride(ilvl=0).add_startOverride(start)
        return new_num.numId

    @staticmethod
    def _set_numbering(paragraph, num_id: int) -> None:
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num_id

    def _list(self, items: typing.Sequence[list[ir.Block]], ordered: bool, base) -> None:
        style = "List Number" if ordered else "List Bullet"
        num_id = None
        for item in items:
            if not item:
                continue
            first = item[0]
            if isinstance(first, ir.Paragraph):
                p = base.add_paragraph(style=style)
                if ordered:
                    if num_id is None:
                        num_id = self._new_numbering(p)
                    self._set_numbering(p, num_id)
                self._inlines(p, first.children)
                self._render_blocks(item[1:], base)
            else:
                p = base.add_paragraph(style=style)
                if ordered:
                    if num_id is None:
                        num_id = self._new_numbering(p)
                    self._set_numbering(p, num_id)
                self._render_blocks(item, base)

    def _inlines(self, paragraph, inlines: typing.Sequence[ir.Inline]) -> None:
        self._render_inlines(paragraph=paragraph, parent=paragraph._p, inlines=inlines)

    def _render_inlines(
            self, paragraph, parent, inlines: typing.Sequence[ir.Inline],
            *,
            bold: bool = False, italic: bool = False, character_style: str | None = None,
    ) -> None:
        for inline in inlines:
            if isinstance(inline, ir.Text):
                self._add_text_run(parent, inline.text, bold=bold, italic=italic, character_style=character_style)
            elif isinstance(inline, ir.Code):
                self._add_text_run(parent, inline.text, bold=bold, italic=italic, character_style=character_style, font="Courier New")
            elif isinstance(inline, ir.Strong):
                self._render_inlines(paragraph, parent, inline.children, bold=True, italic=italic, character_style=character_style)
            elif isinstance(inline, ir.Emphasis):
                self._render_inlines(paragraph, parent, inline.children, bold=bold, italic=True, character_style=character_style)
            elif isinstance(inline, ir.Link):
                if inline.destination.startswith("asn1-ref:"):
                    target = inline.destination.removeprefix("asn1-ref:")
                    self._add_internal_hyperlink(paragraph, parent, inline.children, self._asn1_bookmark_id(target), bold=bold, italic=italic)
                else:
                    self._add_hyperlink(paragraph, parent, inline.children, inline.destination, bold=bold, italic=italic)
            elif isinstance(inline, ir.SoftBreak):
                self._add_text_run(parent," ", bold=bold, italic=italic, character_style=character_style)
            elif isinstance(inline, ir.HardBreak):
                self._add_break_run(parent, bold=bold, italic=italic, character_style=character_style)
            elif isinstance(inline, ir.ASN1Reference):
                add_internal_hyperlink(parent, inline.name, self._asn1_bookmark_id(inline.ref_id))
            else:
                raise NotImplementedError(type(inline).__name__)

    def _add_text_run(self, parent, text: str, *, bold: bool = False, italic: bool = False, character_style: str | None = None, font: str | None = None) -> None:
        run = docx.oxml.OxmlElement("w:r")
        r_pr = self._run_properties(bold=bold, italic=italic, character_style=character_style, font=font)
        if r_pr is not None:
            run.append(r_pr)
        t = docx.oxml.OxmlElement("w:t")
        if text[:1].isspace() or text[-1:].isspace():
            t.set(docx.oxml.ns.qn("xml:space"), "preserve")
        t.text = text
        run.append(t)
        parent.append(run)

    def _add_break_run(self, parent, *, bold: bool = False, italic: bool = False, character_style: str | None = None) -> None:
        run = docx.oxml.OxmlElement("w:r")
        r_pr = self._run_properties(bold=bold, italic=italic, character_style=character_style)
        if r_pr is not None:
            run.append(r_pr)
        run.append(docx.oxml.OxmlElement("w:br"))
        parent.append(run)

    @staticmethod
    def _run_properties(
            *, bold: bool = False, italic: bool = False,character_style: str | None = None, font: str | None = None
    ):
        if not any((bold, italic, character_style, font)):
            return None
        r_pr = docx.oxml.OxmlElement("w:rPr")
        if character_style:
            style = docx.oxml.OxmlElement("w:rStyle")
            style.set(docx.oxml.ns.qn("w:val"), character_style)
            r_pr.append(style)
        if bold:
            r_pr.append(docx.oxml.OxmlElement("w:b"))
        if italic:
            r_pr.append(docx.oxml.OxmlElement("w:i"))
        if font:
            fonts = docx.oxml.OxmlElement("w:rFonts")
            fonts.set(docx.oxml.ns.qn("w:ascii"), font)
            fonts.set(docx.oxml.ns.qn("w:hAnsi"), font)
            r_pr.append(fonts)
        return r_pr

    def _add_hyperlink(
            self, paragraph, parent, children: typing.Sequence[ir.Inline],
            url: str, *, bold: bool = False, italic: bool = False,
    ) -> None:
        r_id = paragraph.part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = docx.oxml.OxmlElement("w:hyperlink")
        hyperlink.set(docx.oxml.ns.qn("r:id"), r_id)
        self._render_inlines(paragraph, hyperlink, children, bold=bold, italic=italic, character_style="Hyperlink")
        parent.append(hyperlink)

    def _add_internal_hyperlink(
            self, paragraph, parent, children: typing.Sequence[ir.Inline],
            bookmark: str, *, bold: bool = False, italic: bool = False,
    ) -> None:
        hyperlink = docx.oxml.OxmlElement("w:hyperlink")
        hyperlink.set(docx.oxml.ns.qn("w:anchor"), bookmark)
        self._render_inlines(paragraph, hyperlink, children, bold=bold, italic=italic, character_style="Hyperlink")
        parent.append(hyperlink)

    def _asn1_type(self, element: ir.ASN1Type, base) -> None:
        if not self.asn1:
            raise ValueError("ASN.1 types not loaded")
        try:
            definition = self.asn1["assignments"][element.ref]
        except KeyError:
            raise ValueError(f"Unknown ASN.1 type: {element.ref}")

        if definition["markdown_doc"]:
            self._render_blocks(definition["markdown_doc"], base)

        if definition["assignment"]["assignment_type"] == "type":
            self._render_asn1_type(definition["name"], definition["assignment"]["type"], base)
        if definition["assignment"]["assignment_type"] == "value":
            self._render_asn1_value_assignment(definition["name"], definition["assignment"], base)
        elif definition["assignment"]["assignment_type"] == "object_set":
            self._render_asn1_object_set(definition["name"], definition["assignment"]["class_ref"], definition["assignment"]["set"], base)
        elif definition["assignment"]["assignment_type"] == "object_class":
            self._render_asn1_object_class(definition["name"], definition["assignment"]["class"], base)

    def _render_asn1_type(self, name: str, type_def: dict, base) -> None:
        kind = type_def.get("type")
        if kind == "sequence":
            self._render_asn1_sequence(name, type_def, base)
        elif kind == "choice":
            self._render_asn1_choice(name, type_def, base)
        elif kind == "enum":
            self._render_asn1_enum(name, type_def, base)

    def _render_asn1_value_assignment(self, name: str, value_def: dict, base) -> None:
        kind = value_def["type"].get("type")
        if kind == "object_identifier":
            self._render_asn1_oid(name, value_def, base)

    def _render_asn1_sequence(self, name: str, type_def, base) -> None:
        p = base.add_paragraph(style="Table header")
        p.add_run(f"SEQUENCE: {name}")
        add_bookmark(p, self._asn1_bookmark_id(type_def["id"]), self.bookmark_id)
        self.bookmark_id += 1

        table = base.add_table(rows=1, cols=4)
        table.style = "Table Grid"

        headers = ["Member", "Type", "Presence", "Description"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True

        for component_name, component in type_def["components"].items():
            row = table.add_row()
            child_type = component.get("type", {})

            row.cells[0].text = component_name
            row.cells[0].paragraphs[0].style = "Code"

            self._render_blocks([ir.Paragraph(self._asn1_type_name(child_type))], row.cells[1])
            e = row.cells[1].paragraphs[0]._element
            e.getparent().remove(e)

            if component["optional"]:
                row.cells[2].text = "O"
            else:
                row.cells[2].text = "M"
            row.cells[2].paragraphs[0].runs[0].bold = True

            if component["markdown_doc"]:
                e = row.cells[3].paragraphs[0]._element
                e.getparent().remove(e)
                self._render_blocks(component["markdown_doc"], row.cells[3])
            if constraint_blocks := self._asn1_constraint_description(child_type):
                self._render_blocks(constraint_blocks, row.cells[3])

            for cell in row.cells:
                cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.TOP

        set_repeat_table_header(table.rows[0])
        make_table_autosize(table)

        for component_name, component in type_def["components"].items():
            ct = component["type"]["type"]
            if ct in {"sequence", "choice", "set", "enum"}:
                self._render_asn1_type(f"{name}.{component_name}", component["type"], base)
            elif ct in {"sequence_of", "set_of"}:
                if component["type"]["inner_type"]["type"] in {"sequence", "choice", "set", "enum"}:
                    self._render_asn1_type(f"{name}.{component_name}", component["type"]["inner_type"], base)

    def _render_asn1_choice(self, name: str, type_def, base) -> None:
        p = base.add_paragraph(style="Table header")
        p.add_run(f"CHOICE: {name}")
        add_bookmark(p, self._asn1_bookmark_id(type_def["id"]), self.bookmark_id)
        self.bookmark_id += 1

        table = base.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        headers = ["Variant", "Type", "Description"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True

        for component_name, component in type_def["components"].items():
            row = table.add_row()
            child_type = component.get("type", {})

            row.cells[0].text = component_name
            row.cells[0].paragraphs[0].style = "Code"

            self._render_blocks([ir.Paragraph(self._asn1_type_name(child_type))], row.cells[1])
            e = row.cells[1].paragraphs[0]._element
            e.getparent().remove(e)

            if component["markdown_doc"]:
                e = row.cells[2].paragraphs[0]._element
                e.getparent().remove(e)
                self._render_blocks(component["markdown_doc"], row.cells[2])
            if constraint_blocks := self._asn1_constraint_description(child_type):
                self._render_blocks(constraint_blocks, row.cells[2])

            for cell in row.cells:
                cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.TOP

        set_repeat_table_header(table.rows[0])
        make_table_autosize(table)

        for component_name, component in type_def["components"].items():
            ct = component["type"]["type"]
            if ct in {"sequence", "choice", "set", "enum"}:
                self._render_asn1_type(f"{name}.{component_name}", component["type"], base)
            elif ct in {"sequence_of", "set_of"}:
                if component["type"]["inner_type"]["type"] in {"sequence", "choice", "set", "enum"}:
                    self._render_asn1_type(f"{name}.{component_name}", component["type"]["inner_type"], base)

    def _render_asn1_enum(self, name: str, type_def, base) -> None:
        p = base.add_paragraph(style="Table header")
        p.add_run(f"ENUMERATION: {name}")
        add_bookmark(p, self._asn1_bookmark_id(type_def["id"]), self.bookmark_id)
        self.bookmark_id += 1

        table = base.add_table(rows=1, cols=3)
        table.style = "Table Grid"

        headers = ["Name", "Value", "Description"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True

        for value in type_def["values"]:
            row = table.add_row()

            row.cells[0].text = value["name"]
            row.cells[0].paragraphs[0].style = "Code"

            row.cells[1].text = str(value.get("value", ""))

            if value["markdown_doc"]:
                e = row.cells[2].paragraphs[0]._element
                e.getparent().remove(e)
                self._render_blocks(value["markdown_doc"], row.cells[2])

            for cell in row.cells:
                cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.TOP

        set_repeat_table_header(table.rows[0])
        make_table_autosize(table)

    def _asn1_type_name(self, type_def: dict):
        kind = type_def.get("type")
        if kind == "character_string":
            return [ir.Code(type_def.get("string_type", "CharacterString"))]
        elif kind == "octet_string":
            return [ir.Code("OCTET STRING")]
        elif kind == "bit_string":
            return [ir.Code("BIT STRING")]
        elif kind == "integer":
            return [ir.Code("INTEGER")]
        elif kind == "boolean":
            return [ir.Code("BOOLEAN")]
        elif kind == "null":
            return [ir.Code("NULL")]
        elif kind == "object_identifier":
            return [ir.Code("OBJECT IDENTIFIER")]
        elif kind == "sequence":
            return [ir.ASN1Reference("SEQUENCE", type_def["id"])]
        elif kind == "choice":
            return [ir.ASN1Reference("CHOICE", type_def["id"])]
        elif kind == "enumerated":
            return [ir.Code("ENUMERATED")]
        elif kind == "sequence_of":
            return [ir.Code("SEQUENCE OF ")] + self._asn1_type_name(type_def["inner_type"])
        elif kind == "set_of":
            return [ir.Code("SET OF ")] + self._asn1_type_name(type_def["inner_type"])
        elif kind == "reference":
            return [ir.ASN1Reference(type_def["name"], type_def["ref_id"])]
        else:
            raise NotImplementedError(f"Type {kind}")

    def _asn1_value(self, value: dict):
        kind = value.get("value_type")
        if kind == "integer":
            return [ir.Text(str(value["value"]))]
        elif kind == "character_string":
            return [ir.Text(str(value["value"]))]
        elif kind == "boolean":
            return [ir.Text("TRUE" if value["value"] else "FALSE")]
        elif kind == "choice":
            return [
                ir.Text(value["variant"]),
                ir.Text("("),
                *self._asn1_value(value["value"]),
                ir.Text(")"),
            ]
        elif kind == "sequence":
            o = [ir.Text("{")]
            for index, field in enumerate(value.get("fields", [])):
                if index:
                    o.append(ir.Text(", "))
                o.append(ir.Text(field["name"]))
                o.append(ir.Text(" = "))
                o.extend(self._asn1_value(field["value"]))
            o.append(ir.Text("}"))
            return o
        elif kind == "reference":
            return [ir.ASN1Reference(value["name"], value["ref_id"])]
        else:
            raise NotImplementedError(f"ASN.1 value type {kind} is not supported")

    def _render_asn1_oid(self, name: str, value_def: dict, base) -> None:
        p = base.add_paragraph(style="Table header")
        p.add_run(f"OBJECT IDENTIFIER: {name}")
        add_bookmark(p, self._asn1_bookmark_id(value_def["type"]["id"]), self.bookmark_id)
        self.bookmark_id += 1

        table = base.add_table(rows=0, cols=2)
        table.style = "Table Grid"

        row = table.add_row()

        row.cells[0].text = "ASN.1"
        row.cells[0].paragraphs[0].runs[0].bold = True
        p = row.cells[1].paragraphs[0]
        p.clear()
        self._add_text_run(p._p," / ".join(
            self._oid_component_symbolic(component)
            for component in value_def["value"]["components"]
        ), font="Courier New")

        for cell in row.cells:
            cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.TOP

        row = table.add_row()

        row.cells[0].text = "Numeric"
        row.cells[0].paragraphs[0].runs[0].bold = True
        p = row.cells[1].paragraphs[0]
        p.clear()
        self._add_text_run(p._p, value_def["value"]["numeric_oid"], font="Courier New")

        for cell in row.cells:
            cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.TOP

        make_table_autosize(table)

    @staticmethod
    def _oid_component_symbolic(component: dict) -> str:
        name = component.get("name")
        number = component["number"]

        if name is not None:
            return f"{name}({number})"

        return str(number)

    def _render_asn1_object_class(self, name: str, class_def, base) -> None:
        p = base.add_paragraph(style="Table header")
        p.add_run(f"OBJECT CLASS: {name}")
        add_bookmark(p, self._asn1_bookmark_id(class_def["id"]), self.bookmark_id)
        self.bookmark_id += 1

        table = base.add_table(rows=1, cols=4)
        table.style = "Table Grid"

        headers = ["Field", "Type", "Presence", "Description"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True

        for field_name, field in class_def["fields"].items():
            row = table.add_row()

            row.cells[0].text = field_name[1:]
            row.cells[0].paragraphs[0].style = "Code"

            if field["field_type"] == "type":
                self._render_blocks([ir.Paragraph([ir.Text("Any type.")])], row.cells[1])
            elif field["field_type"] == "fixed_type_value":
                self._render_blocks([ir.Paragraph(self._asn1_type_name(field["type"]))], row.cells[1])

            e = row.cells[1].paragraphs[0]._element
            e.getparent().remove(e)

            if field["optional"]:
                row.cells[2].text = "O"
            else:
                row.cells[2].text = "M"
            row.cells[2].paragraphs[0].runs[0].bold = True

            if field["markdown_doc"]:
                self._render_blocks(field["markdown_doc"], row.cells[3])
            if field["field_type"] == "fixed_type_value":
                if constraint_blocks := self._asn1_constraint_description(field["type"]):
                    self._render_blocks(constraint_blocks, row.cells[3])
            if field.get("unique", False):
                self._render_blocks([ir.Paragraph([
                    ir.Text("Values of this field "),
                    ir.Strong([ir.Text("SHALL")]),
                    ir.Text(" be unique within an object set.")
                ])], row.cells[3])
            if len(row.cells[3].paragraphs) > 1:
                e = row.cells[3].paragraphs[0]._element
                e.getparent().remove(e)

            for cell in row.cells:
                cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.TOP

        set_repeat_table_header(table.rows[0])
        make_table_autosize(table)

        for field_name, field in class_def["fields"].items():
            if field["field_type"] == "fixed_type_value" and field["type"].get("type") in {"sequence", "choice", "set", "enum"}:
                self._render_asn1_type(f"{name}.{field_name}", field["type"], base)

    def _render_asn1_object_set(self, name: str, class_ref: dict, object_set_def: dict, base) -> None:
        members = object_set_def.get("members", [])
        p = base.add_paragraph(style="Table header")
        p.add_run(f"OBJECT SET: {name}, of OBJECT CLASS ")
        add_internal_hyperlink(p._p, class_ref["name"], self._asn1_bookmark_id(class_ref["id"]))
        add_bookmark(p, self._asn1_bookmark_id(object_set_def["id"]), self.bookmark_id)
        self.bookmark_id += 1

        columns: list[tuple[str, tuple[str, ...]]] = []
        rows: list[dict[tuple[str, ...], dict | None]] = []

        for member in members:
            flattened: dict[tuple[str, ...], dict | None] = {}
            for field in member.get("object", {}).get("fields", []):
                flattened.update(self._flatten_asn1_object_value((field["name"][1:],), field.get("value")))
            rows.append(flattened)
            for path in flattened:
                if path not in [column_path for _, column_path in columns]:
                    columns.append((".".join(path), path))

        table = base.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"

        for cell, (heading, _) in zip(table.rows[0].cells, columns):
            cell.text = heading
            cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True

        for member, flattened in zip(members, rows):
            row = table.add_row()
            for cell, (_, path) in zip(row.cells, columns):
                value = flattened.get(path)
                if value is not None:
                    if "value_type" in value:
                        self._render_inlines(cell.paragraphs[0], cell.paragraphs[0]._p, self._asn1_value(value))
                    elif "type" in value:
                        self._render_inlines(cell.paragraphs[0], cell.paragraphs[0]._p, self._asn1_type_name(value))
                cell.paragraphs[0].style = "Code"
                cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.TOP

            if member.get("markdown_doc"):
                desc_row = table.add_row()
                desc_cell = desc_row.cells[0]
                for cell in desc_row.cells[1:]:
                    desc_cell = desc_cell.merge(cell)
                desc_cell.vertical_alignment = docx.enum.table.WD_CELL_VERTICAL_ALIGNMENT.TOP
                self._render_blocks(member["markdown_doc"], desc_cell)
                e = desc_cell.paragraphs[0]._element
                e.getparent().remove(e)

        set_repeat_table_header(table.rows[0])
        make_table_autosize(table)

    def _flatten_asn1_object_value(self, path: tuple[str, ...], value: dict | None) -> dict[tuple[str, ...], dict | None]:
        if value is None:
            return {path: None}

        kind = value.get("value_type")

        if kind == "sequence":
            result: dict[tuple[str, ...], dict | None] = {}
            for field in value.get("fields", []):
                field_path = path + (field["name"],)
                result.update(self._flatten_asn1_object_value(field_path, field.get("value")))
            return result

        return {path: value}

    def _asn1_bookmark_id(self, type_id):
        if type_id in self.asn1_bookmarks:
            return self.asn1_bookmarks[type_id]
        else:
            b = f"asn1_{self.asn1_bookmark_id}"
            self.asn1_bookmarks[type_id] = b
            self.asn1_bookmark_id += 1
            return b

    def _asn1_constraint_description(self, type_def: dict) -> list[ir.Block]:
        constraint = type_def.get("constraint")
        if constraint is None:
            return []

        descriptions = self._asn1_constraint_sentences(constraint, type_def=type_def)
        if not descriptions:
            return []

        return [ir.Paragraph(children=join_inline_sentences(descriptions))]

    def _asn1_constraint_sentences(self, constraint: dict, *, type_def: dict) -> list[list[ir.Inline]]:
        kind = constraint["constraint_type"]

        if kind == "intersection":
            result = []
            for child in constraint.get("constraints", []):
                result.extend(self._asn1_constraint_sentences(child, type_def=type_def))
            return result
        elif kind == "union":
            return self._asn1_union_constraint(constraint)
        elif kind == "inverse":
            return self._asn1_inverse_constraint(constraint)
        elif kind == "size_fixed":
            return [self._asn1_size_fixed_constraint(constraint, type_def=type_def)]
        elif kind == "size_range":
            return [self._asn1_size_range_constraint(constraint, type_def=type_def)]
        elif kind == "single_value":
            return [self._asn1_single_value_constraint(constraint)]
        elif kind == "value_range":
            return [self._asn1_value_range_constraint(constraint)]
        elif kind == "permitted_alphabet":
            return [self._asn1_alphabet_constraint(constraint)]
        elif kind == "table":
            return [self._asn1_table_constraint(constraint)]

        raise NotImplementedError(f"ASN.1 constraint {kind} is not supported")

    def _asn1_value_text(self, value: dict) -> str:
        kind = value["value_type"]
        if kind == "integer":
            return str(value["value"])
        elif kind == "character_string":
            return str(value["value"])
        elif kind == "boolean":
            return "TRUE" if value["value"] else "FALSE"
        raise NotImplementedError(f"ASN.1 value {kind} is not supported")

    def _asn1_value_inline(self, value: dict) -> ir.Inline:
        return ir.Code(self._asn1_value_text(value))

    def _asn1_size_unit(self, type_def: dict, count: int | None = None) -> str:
        kind = type_def.get("type")
        if kind == "character_string":
            singular, plural = "character", "characters"
        elif kind == "octet_string":
            singular, plural = "octet", "octets"
        elif kind == "bit_string":
            singular, plural = "bit", "bits"
        elif kind in {"sequence_of", "set_of"}:
            singular, plural = "element", "elements"
        else:
            raise NotImplementedError(f"SIZE description for ASN.1 type {kind}")
        return singular if count == 1 else plural

    def _asn1_size_range_constraint(self, constraint: dict, *, type_def: dict) -> list[ir.Inline]:
        minimum = int(self._asn1_value_text(constraint["min_size"]))
        maximum = int(self._asn1_value_text(constraint["max_size"]))
        unit = self._asn1_size_unit(type_def)

        minimum_inclusive = "inclusive" if constraint["min_inclusive"] else "exclusive"
        maximum_inclusive = "inclusive" if constraint["max_inclusive"] else "exclusive"
        return [
            ir.Text("The value "),
            ir.Strong([ir.Text("SHALL")]),
            ir.Text(f" consist of between {minimum} ({minimum_inclusive}) and {maximum} ({maximum_inclusive}) {unit}"),
        ]

    def _asn1_size_fixed_constraint(self, constraint: dict, *, type_def: dict) -> list[ir.Inline]:
        size = int(self._asn1_value_text(constraint["size"]))
        unit = self._asn1_size_unit(type_def, count=size)
        return [
            ir.Text("The value "),
            ir.Strong([ir.Text("SHALL")]),
            ir.Text(f" consist of exactly {size} {unit}"),
        ]

    def _asn1_value_range_constraint(self, constraint: dict) -> list[ir.Inline]:
        minimum = self._asn1_value_text(constraint["min"])
        minimum_inclusive = "inclusive" if constraint["min_inclusive"] else "exclusive"
        maximum = self._asn1_value_text(constraint["max"])
        maximum_inclusive = "inclusive" if constraint["max_inclusive"] else "exclusive"
        return [
            ir.Text("The value "),
            ir.Strong([ir.Text("SHALL")]),
            ir.Text(f" be between {minimum} ({minimum_inclusive}) and {maximum} ({maximum_inclusive})"),
        ]

    def _asn1_single_value_constraint(self, constraint: dict) -> list[ir.Inline]:
        value = self._asn1_value_text(constraint["value"])
        return [
            ir.Text("The value "),
            ir.Strong([ir.Text("SHALL")]),
            ir.Text(f" be {value}"),
        ]

    def _asn1_inverse_constraint(self, constraint: dict) -> list[list[ir.Inline]]:
        exclusion = constraint["exclusion"]

        if exclusion["constraint_type"] == "union":
            children = exclusion["constraints"]
            if all(child["constraint_type"] == "single_value" for child in children):
                values = [child["value"] for child in children]
                return [[
                    ir.Text("The value "),
                    ir.Strong([ir.Text("SHALL NOT")]),
                    ir.Text(" be "),
                    *inline_list([self._asn1_value_inline(v) for v in values]),
                ]]

        raise NotImplementedError("Unsupported ASN.1 inverse constraint")

    def _asn1_table_constraint(self, constraint: dict) -> list[ir.Inline]:
        ref = constraint["object_set_ref"]
        return [
            ir.Text("The value "),
            ir.Strong([ir.Text("SHALL")]),
            ir.Text(" be a member of "),
            ir.ASN1Reference(ref["name"], ref["id"]),
        ]

    def _asn1_alphabet_constraint(self, constraint: dict) -> list[ir.Inline]:
        alphabet = constraint["alphabet"]
        return [
            ir.Text("Each character "),
            ir.Strong([ir.Text("SHALL")]),
            ir.Text(" be "),
            *self._asn1_alphabet_description(alphabet),
        ]

    def _asn1_alphabet_description(self, constraint: dict) -> list[ir.Inline]:
        kind = constraint["constraint_type"]
        if kind == "single_value":
            return [ir.Code(self._asn1_value_text(constraint["value"]))]
        elif kind == "value_range":
            minimum = self._asn1_value_inline(constraint["min"])
            minimum_inclusive = "inclusive" if constraint["min_inclusive"] else "exclusive"
            maximum = self._asn1_value_inline(constraint["max"])
            maximum_inclusive = "inclusive" if constraint["max_inclusive"] else "exclusive"

            return [
                ir.Text("in the range "), minimum, ir.Text(f" ({minimum_inclusive}) to "),
                maximum, ir.Text(f" ({maximum_inclusive})"),
            ]
        elif kind == "union":
            alternatives = [
                self._asn1_alphabet_description(child)
                for child in constraint["constraints"]
            ]
            return join_inline_alternatives(alternatives)
        elif kind == "intersection":
            constraints = [
                self._asn1_alphabet_description(child)
                for child in constraint["constraints"]
            ]
            result: list[ir.Inline] = []
            for i, child in enumerate(constraints):
                if i:
                    result.append(ir.Text(" and "))
                result.extend(child)
            return result
        elif kind == "inverse":
            return [
                ir.Text("other than "),
                *self._asn1_alphabet_description(constraint["exclusion"]),
            ]

        raise NotImplementedError(
            f"Unsupported permitted alphabet constraint: {kind!r}"
        )


def inline_list(values: list[ir.Inline]) -> list[ir.Inline]:
    result = []
    for i, value in enumerate(values):
        if i != 0:
            if i == len(values) - 1:
                result.append(
                    ir.Text(" or " if len(values) == 2 else ", or ")
                )
            else:
                result.append(ir.Text(", "))
        result.append(value)
    return result

def join_inline_sentences(sentences: list[list[ir.Inline]]) -> list[ir.Inline]:
    result = []
    for sentence in sentences:
        if result:
            result.append(ir.Text(" "))
        result.extend(sentence)
        result.append(ir.Text("."))
    return result

def join_inline_alternatives(alternatives: list[list[ir.Inline]]) -> list[ir.Inline]:
    result: list[ir.Inline] = []
    for i, alternative in enumerate(alternatives):
        if i:
            if i == len(alternatives) - 1:
                result.append(ir.Text(" or " if len(alternatives) == 2 else ", or "))
            else:
                result.append(ir.Text(", "))
        result.extend(alternative)
    return result

def ensure_style(doc: docx.document.Document, name: str, base: str = "Normal"):
    styles = doc.styles
    try:
        return styles[name]
    except KeyError:
        from docx.enum.style import WD_STYLE_TYPE
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        try:
            style.base_style = styles[base]
        except KeyError:
            pass
        return style


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    safe = re.sub(r"[^A-Za-z0-9_]", "_", name)
    if not safe or not safe[0].isalpha():
        safe = "b_" + safe

    start = docx.oxml.OxmlElement("w:bookmarkStart")
    start.set(docx.oxml.ns.qn("w:id"), str(bookmark_id))
    start.set(docx.oxml.ns.qn("w:name"), safe)

    end = docx.oxml.OxmlElement("w:bookmarkEnd")
    end.set(docx.oxml.ns.qn("w:id"), str(bookmark_id))

    paragraph._p.insert(0, start)
    paragraph._p.insert(1, end)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE, is_external=True)

    hyperlink = docx.oxml.OxmlElement("w:hyperlink")
    hyperlink.set(docx.oxml.ns.qn("r:id"), r_id)

    run = docx.oxml.OxmlElement("w:r")
    r_pr = docx.oxml.OxmlElement("w:rPr")

    r_style = docx.oxml.OxmlElement("w:rStyle")
    r_style.set(docx.oxml.ns.qn("w:val"), "Hyperlink")
    r_pr.append(r_style)

    run.append(r_pr)

    text_element = docx.oxml.OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_internal_hyperlink(parent, text: str, bookmark: str) -> None:
    hyperlink = docx.oxml.OxmlElement("w:hyperlink")
    hyperlink.set(docx.oxml.ns.qn("w:anchor"), bookmark)

    run = docx.oxml.OxmlElement("w:r")
    r_pr = docx.oxml.OxmlElement("w:rPr")

    r_style = docx.oxml.OxmlElement("w:rStyle")
    r_style.set(docx.oxml.ns.qn("w:val"), "Hyperlink")
    r_pr.append(r_style)

    run.append(r_pr)

    text_element = docx.oxml.OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)

    hyperlink.append(run)
    parent.append(hyperlink)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = docx.oxml.OxmlElement("w:tblHeader")
    tbl_header.set(docx.oxml.ns.qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(run, instruction, helper_text: str = ""):
    begin = docx.oxml.OxmlElement("w:fldChar", attrs={
        docx.oxml.ns.qn("w:fldCharType"): "begin"
    })
    instr = docx.oxml.OxmlElement("w:instrText", attrs={
        docx.oxml.ns.qn("xml:space"): "preserve"
    })
    instr.text = instruction
    sep = docx.oxml.OxmlElement("w:fldChar", attrs={
        docx.oxml.ns.qn("w:fldCharType"): "separate"
    })
    text = docx.oxml.OxmlElement("w:t")
    text.text = helper_text
    end = docx.oxml.OxmlElement("w:fldChar", attrs={
        docx.oxml.ns.qn("w:fldCharType"): "end"
    })
    run._r.extend([begin, instr, sep, text, end])


def make_table_autosize(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    layout = tbl_pr.find(docx.oxml.ns.qn("w:tblLayout"))
    if layout is None:
        layout = docx.oxml.OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(docx.oxml.ns.qn("w:type"), "autofit")

    tbl_w = tbl_pr.find(docx.oxml.ns.qn("w:tblW"))
    if tbl_w is None:
        tbl_w = docx.oxml.OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)

    tbl_w.set(docx.oxml.ns.qn("w:type"), "pct")
    tbl_w.set(docx.oxml.ns.qn("w:w"), "5000")

    for grid_col in tbl.tblGrid.findall(docx.oxml.ns.qn("w:gridCol")):
        grid_col.attrib.pop(docx.oxml.ns.qn("w:w"), None)

    for row in table.rows:
        for cell in row.cells:
            tc_w = cell._tc.get_or_add_tcPr().find(docx.oxml.ns.qn("w:tcW"))
            if tc_w is not None:
                cell._tc.get_or_add_tcPr().remove(tc_w)